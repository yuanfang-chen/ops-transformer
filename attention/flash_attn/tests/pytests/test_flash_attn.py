import os
import argparse
import openpyxl
import torch
import torch_npu
import torch.nn.functional as f
import math
import numpy as np
from cpu_impl import tforward
from test_utils import generate_qkv, generate_pse, generate_npu_mask, trans_bnsd_to_layout, get_seqlen_list
from npu_impl import flash_attn_npu

from npu_ops_transformer.ops import npu_flash_attn_meta

def check_result(expect, result, test_name):
    print(f"开始比对{test_name}的精度.")
    ratio_threshold = 0.005
    threshold_diff = 0.005
    if expect.shape == result.shape:
        if torch.all(torch.eq(expect, result)):
            print(f"{test_name} 计算结果完全一致.")
            ratio_diff = 0
            diff = 0
            return 1, 0, 0
        else:
            diff = torch.abs(expect.sub(result))
            min = torch.empty(expect.shape, dtype=torch.float16)
            min.fill_(0.000025)
            max = torch.max(torch.abs(expect), torch.abs(result))
            threshold = torch.max(max.mul(ratio_threshold), min)
            mask = diff > threshold
            num_diff = torch.sum(mask)
            ratio_diff = num_diff / torch.numel(expect)
            if ratio_diff > threshold_diff:
                print(f"warning: {test_name} 计算结果有{num_diff}个元素的偏差超过阈值{ratio_threshold:.2%},占比为{ratio_diff:.2%}!")
            else:
                print(f"info: {test_name} 计算结果有{num_diff}个元素的偏差超过阈值{ratio_threshold:.2%},占比为{ratio_diff:.2%}.")
            ratio = (1 - ratio_diff).cpu().detach().numpy()
            max = torch.max(diff).cpu().detach().numpy()
            print(f"{test_name} diff_max: {max:.8f}")
            sum = torch.sum(diff).cpu().detach().numpy()
            print(f"{test_name} diff_sum: {sum:.8f}")
    else:
        print(f"error: {test_name} 计算结果错误,shape与标杆不匹配!")
        print(f"error: expect shape: {expect.shape}")
        print(f"error: result shape: {result.shape}")
        ratio = 0
        max = 999999
        sum = 999999
    return ratio, max, sum


def call_flash_attn(test_name, **kwargs):
    b = kwargs.get("B", 1)
    n1 = kwargs.get("N1")
    n2 = kwargs.get("N2", n1)
    sq = kwargs.get("S1", -1)
    skv = kwargs.get("S2", sq)
    d = kwargs.get("D")
    d_v = kwargs.get("DV", d)
    d_rope = kwargs.get("DRope", 0)
    input_layout = kwargs.get("input_layout")
    scale = kwargs.get("scale", 1 / (d ** 0.5))
    pse_type = int(kwargs.get("pse_type") if (kwargs.get("pse_type") != '') else 0)
    pse_layout = kwargs.get("pse_layout", "none").lower()
    q_start_idx = kwargs.get("q_start_idx", 0)
    kv_start_idx = kwargs.get("kv_start_idx", 0)
    dtype = kwargs.get("Dtype", torch.bfloat16)
    if dtype == 'fp16':
        pttype = torch.float16
        input_dtype = torch.float16
    if dtype == 'bf16':
        pttype = torch.bfloat16
        input_dtype = torch.bfloat16

    sparse_mode = kwargs.get("sparse_mode", None)
    pre_tokens = kwargs.get("pre_tokens", 2147483647)
    next_tokens = kwargs.get("next_tokens", 2147483647)
    prefix = kwargs.get("prefix", [])
    pse_b = b
    pse_s1 = sq
    pse_s2 = skv
    actual_seq_qlen = None
    actual_seq_kvlen = None
    if input_layout == "TND":
        actual_seq_qlen = list(kwargs.get("seqlens_list_q"))
        actual_seq_kvlen = list(kwargs.get("seqlens_list_kv", actual_seq_qlen))
        seqlen_q_list = get_seqlen_list(actual_seq_qlen)
        seqlen_k_list = get_seqlen_list(actual_seq_kvlen)
        sq = max(actual_seq_qlen)
        skv = max(actual_seq_kvlen)
        pse_b = len(actual_seq_qlen)
        pse_s1 = seqlen_q_list.max()
        pse_s2 = seqlen_k_list.max()
        if pse_layout in ["bnhs", "1nhs"]:
            pse_s1 = max(1024, pse_s1)

    pse_cpu, pse_npu = generate_pse(pse_b, n1, pse_s1, pse_s2, pse_type, pse_layout, pttype, q_start_idx, kv_start_idx)
    q, k, v, q_rope, k_rope, qf, kf = generate_qkv(b, n1, n2, sq, skv, d, d_v, d_rope, input_layout, input_dtype)
    out, x_max, x_sum = tforward(qf, kf, v, pse_cpu, **kwargs)

    atten_mask = generate_npu_mask(b, sq, skv, sparse_mode, pre_tokens, next_tokens, prefix)
    npu_out, npu_max, npu_sum = flash_attn_npu(q, k, v, q_rope, k_rope, atten_mask, pse_npu, **kwargs)

    out = trans_bnsd_to_layout(out, input_layout)
    check_result(out.float(), npu_out.float(), "out")
    
    save_dir = "./pytests"
    os.makedirs(save_dir, exist_ok=True)
    save_path = os.path.join(save_dir, f"{test_name}_tensors.pt")
    torch.save({
        'q': q,
        'k': k,
        'v': v,
        'out': out,
        'npu_out': npu_out
    }, save_path)
    print(f"Tensors saved to {save_path}")
    
    # check_result(x_max.contiguous().view(-1), npu_max.contiguous().view(-1), "max")
    # check_result(x_sum.contiguous().view(-1), npu_sum.contiguous().view(-1), "sum")

def get_col_index(table, col_name):
    col_index = None
    for i in range(1,table.max_column + 1):
       if table.cell(row = 1,column=i).value == col_name:
            col_index = i
            break
    return col_index

if __name__ =="__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument('--case_id', type=str, default='all', help='case name')
    parser.add_argument('--device_id', type=int, default='0', help='device_id')
    parser.add_argument('--case_file', type=str, default='FlashAttn', help='case file')
    parser.add_argument('--version', type=str, default='950', help='device version')
    parser.add_argument('--case_level', type=str, default='all',help='running specified level')
    parser.add_argument('--sheet', type=str, default='Sheet1', help='sheet name')
    args = parser.parse_args()
    case_name = args.case_id
    device_id = args.device_id
    case_file = args.case_file
    version = args.version
    sheet = args.sheet

    case_path = f"./excel/{version}/{case_file}.xlsx"
    if os.path.exists(case_path):
        data = openpyxl.load_workbook(case_path)
    else:
        print(f"ERROR: case file {case_path} does not exist.")
        exit()
    # result_xlsx_name = case_file + f"_Result_{datatime.datatime.now().strftime('%Y_%m_%d_%H_%M_%S')}.xlsx"
    table = data[sheet]
    nrows = table.max_row
    Testcase_Name_Col = get_col_index(table, 'Testcase_Name')
    Enable_Col = get_col_index(table, 'Enable')
    case_level_Col = get_col_index(table, 'Level')
    B_col = get_col_index(table, 'B')
    N1_col = get_col_index(table, 'N1')
    N2_col = get_col_index(table, 'N2')
    S1_col = get_col_index(table, 'S1')
    S2_col = get_col_index(table, 'S2')
    seqlens_list_col = get_col_index(table, 'seqlens_list_q')
    seqlens_list_kv_col = get_col_index(table, 'seqlens_list_kv')
    D_col = get_col_index(table, 'D')
    DV_col = get_col_index(table, 'DV')
    DR_col = get_col_index(table, 'DRope')
    Dtype_col = get_col_index(table, 'Dtype')
    outDtype_col = get_col_index(table, 'out_dtype')
    Sparse_col = get_col_index(table, 'sparse_mode')
    prefix_col = get_col_index(table, 'prefix')
    Input_layout_col = get_col_index(table, 'Layout')
    Atten_mask_shape_col = get_col_index(table, 'Atten_mask_Shape')
    Atten_mask_dtype_col = get_col_index(table, 'Atten_mask_Dtype')
    Paddding_mask_col = get_col_index(table, 'Padding_Mask')
    Pse_shape_col = get_col_index(table, 'PSE')
    Pse_mode_col = get_col_index(table, 'PSE_mode')
    Pse_type_col = get_col_index(table, 'pse_type')
    pre_tokens_col = get_col_index(table, 'pre_tokens')
    next_tokens_col = get_col_index(table, 'next_tokens')
    keep_prob_col = get_col_index(table, 'keep_prob')
    q_start_idx_col = get_col_index(table, 'q_start_idx')
    kv_start_idx_col = get_col_index(table, 'kv_start_idx')
    seed_col = get_col_index(table, 'seed')
    offset_col = get_col_index(table, 'offset')
    for i in range(2, nrows + 1):
        case = {}
        case['Testcase_Name'] = table.cell(row=i, column=Testcase_Name_Col).value
        test_name = case['Testcase_Name']
        case['Level'] = table.cell(row=i, column=case_level_Col).value
        case['B'] = int(table.cell(row=i, column=B_col).value)
        case['N1'] = int(table.cell(row=i, column=N1_col).value)
        case['N2'] = int(table.cell(row=i, column=N2_col).value) if table.cell(row=i, column=N2_col).value is not None else case['N1']
        case['S1'] = int(table.cell(row=i, column=S1_col).value)
        case['S2'] = int(table.cell(row=i, column=S2_col).value) if table.cell(row=i, column=S2_col).value is not None else case['S1']
        case['seqlens_list_q'] = eval(table.cell(i, seqlens_list_col).value) if table.cell(i, seqlens_list_col).value is not None else None
        case['seqlens_list_kv'] = case['seqlens_list_q'] if table.cell(i, seqlens_list_col).value is None else eval(table.cell(i, seqlens_list_col).value)
        case['D'] = int(table.cell(row=i, column=D_col).value)
        case['DV'] = int(table.cell(row=i, column=D_col).value) if table.cell(row=i, column=DV_col).value is None else int(table.cell(row=i, column=DV_col).value)
        case['DRope'] = int(table.cell(row=i, column=DR_col).value) if table.cell(row=i, column=DR_col).value is not None else 0
        case['Dtype'] = table.cell(row=i, column=Dtype_col).value
        case['out_dtype'] = 0 if table.cell(row=i, column=outDtype_col).value is None else table.cell(row=i, column=outDtype_col).value
        case['sparse_mode'] = int(table.cell(row=i, column=Sparse_col).value)
        case['prefix'] = eval(table.cell(row=i, column=prefix_col).value) if prefix_col is not None and table.cell(row=i, column=prefix_col).value is None else None
        case['input_layout'] = table.cell(row=i, column=Input_layout_col).value
        case['Atten_mask_Shape'] = table.cell(row=i, column=Atten_mask_shape_col).value
        case['Atten_mask_Dtype'] = table.cell(row=i, column=Atten_mask_dtype_col).value
        case['Padding_Mask'] = table.cell(row=i, column=Paddding_mask_col).value
        case['pse_layout'] = table.cell(row=i, column=Pse_shape_col).value
        case['pse_mode'] = int(table.cell(row=i, column=Pse_mode_col).value) if table.cell(row=i, column=Pse_mode_col).value is not None else 1
        case['pse_type'] = '' if table.cell(row=i, column=Pse_type_col).value is None else table.cell(row=i, column=Pse_type_col).value
        case['pre_tokens'] = int(table.cell(row=i, column=pre_tokens_col).value) if table.cell(row=i, column=pre_tokens_col).value is not None else 65536
        case['next_tokens'] = int(table.cell(row=i, column=next_tokens_col).value) if table.cell(row=i, column=next_tokens_col).value is not None else 65536
        case['keep_prob'] = int(table.cell(row=i, column=keep_prob_col).value) if table.cell(row=i, column=keep_prob_col).value is not None else 1
        case['q_start_idx'] = int(table.cell(row=i, column=q_start_idx_col).value) if table.cell(row=i, column=q_start_idx_col).value is not None else 0
        case['kv_start_idx'] = int(table.cell(row=i, column=kv_start_idx_col).value) if table.cell(row=i, column=kv_start_idx_col).value is not None else 0
        case['seed'] = int(table.cell(row=i, column=seed_col).value) if table.cell(row=i, column=seed_col).value is not None else 0
        case['offset'] = int(table.cell(row=i, column=offset_col).value) if table.cell(row=i, column=offset_col).value is not None else 0

        enable = table.cell(row=i, column=Enable_Col).value
        if (enable.lower() =='enable' and case_name =='all' and case_level == 'all') or case["Testcase_Name"] == case_name or (case["Level"].lower() == case_level.lower() and enable.lower() == 'enable'):
            call_flash_attn(test_name, **case)





